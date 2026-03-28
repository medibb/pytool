#!/usr/bin/env python3
"""
부산대병원 재활의학과 — 프로토콜 제출서 Word(docx) 생성기
가정의학과 프로토콜 제출 견본 형식에 맞춤 (범용)

사용법:
    1. 이 파일 하단의 if __name__ == "__main__" 블록에서 config를 작성
    2. python3 pnuh_protocol_generator.py 실행
    3. 또는 다른 스크립트에서 import 하여 사용:

        from pnuh_protocol_generator import generate_protocol
        generate_protocol(config, output_path)

config 구조:
    {
        "title": "프로토콜 제출서 상단 부제목",
        "header_title": "헤더 테이블에 들어갈 프로토콜명",
        "department": "진료과명",
        "sections": [ ... ],       # 프로토콜 본문 섹션
        "notes": [ ... ],          # 비고 섹션
    }

sections 구조:
    각 section은 dict:
    {
        "label": "섹션 라벨 (왼쪽 병합 셀)",
        "rows": [
            ("소항목명", "BG색상코드", ["내용1", "내용2", ("굵은내용", True), ...]),
            ...
        ]
    }

notes 구조:
    [("텍스트", bold: bool, red: bool), ...]
"""

from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn, nsdecls
from docx.oxml import parse_xml
import os

# ── 표준 색상 ──
COLORS = {
    "HDR":  "2F5496",  # 헤더 (진한 파랑)
    "SEC":  "D6E4F0",  # 섹션 라벨 배경 (연파랑)
    "SUB":  "E2EFDA",  # 일반 소항목 (연녹색)
    "WARN": "FFF2CC",  # 주의 (연노랑)
    "CRIT": "FCE4EC",  # 위험/금기 (연빨강)
    "NOTE": "F2F2F2",  # 비고 (연회색)
}


# ══════════════════════════════════════════════════════════════
# Helper functions
# ══════════════════════════════════════════════════════════════

def _set_cell_shading(cell, color_hex):
    shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{color_hex}"/>')
    cell._tc.get_or_add_tcPr().append(shading_elm)


def _set_cell_text(cell, text, bold=False, size=9, color=None, align=WD_ALIGN_PARAGRAPH.LEFT):
    cell.text = ''
    p = cell.paragraphs[0]
    p.alignment = align
    p.paragraph_format.space_before = Pt(1)
    p.paragraph_format.space_after = Pt(1)
    run = p.add_run(text)
    run.font.size = Pt(size)
    run.font.name = '맑은 고딕'
    run.font.element.rPr.rFonts.set(qn('w:eastAsia'), '맑은 고딕')
    run.bold = bold
    if color:
        run.font.color.rgb = RGBColor(*color)
    return run


def _add_lines(cell, lines, size=8):
    cell.text = ''
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(1)
    p.paragraph_format.space_after = Pt(1)
    for i, line in enumerate(lines):
        if i > 0:
            p.add_run('\n')
        text = line[0] if isinstance(line, tuple) else line
        bold = line[1] if isinstance(line, tuple) else False
        run = p.add_run(text)
        run.font.size = Pt(size)
        run.font.name = '맑은 고딕'
        run.font.element.rPr.rFonts.set(qn('w:eastAsia'), '맑은 고딕')
        run.bold = bold


def _set_table_borders(table):
    tbl = table._tbl
    tblPr = tbl.tblPr if tbl.tblPr is not None else parse_xml(f'<w:tblPr {nsdecls("w")}/>')
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        '<w:top w:val="single" w:sz="6" w:space="0" w:color="000000"/>'
        '<w:left w:val="single" w:sz="6" w:space="0" w:color="000000"/>'
        '<w:bottom w:val="single" w:sz="6" w:space="0" w:color="000000"/>'
        '<w:right w:val="single" w:sz="6" w:space="0" w:color="000000"/>'
        '<w:insideH w:val="single" w:sz="4" w:space="0" w:color="808080"/>'
        '<w:insideV w:val="single" w:sz="4" w:space="0" w:color="808080"/>'
        '</w:tblBorders>'
    )
    tblPr.append(borders)


# ══════════════════════════════════════════════════════════════
# Main generator
# ══════════════════════════════════════════════════════════════

def generate_protocol(config, output_path):
    """
    config dict로부터 프로토콜 제출서 docx를 생성.

    Args:
        config: dict with keys: title, header_title, department, sections, notes
        output_path: 출력 docx 파일 경로
    """
    doc = Document()

    # Page setup
    for section in doc.sections:
        section.top_margin = Cm(1.5)
        section.bottom_margin = Cm(1.5)
        section.left_margin = Cm(1.5)
        section.right_margin = Cm(1.5)

    # Default font
    style = doc.styles['Normal']
    font = style.font
    font.name = '맑은 고딕'
    font.size = Pt(9)
    style.element.rPr.rFonts.set(qn('w:eastAsia'), '맑은 고딕')

    C = COLORS

    # ── Title ──
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run('프로토콜 제출서')
    r.bold = True
    r.font.size = Pt(18)
    r.font.name = '맑은 고딕'
    r.font.element.rPr.rFonts.set(qn('w:eastAsia'), '맑은 고딕')

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(config["title"])
    r.font.size = Pt(11)
    r.font.name = '맑은 고딕'
    r.font.element.rPr.rFonts.set(qn('w:eastAsia'), '맑은 고딕')
    r.font.color.rgb = RGBColor(0x2F, 0x54, 0x96)

    doc.add_paragraph()

    # ── Header table ──
    ht = doc.add_table(rows=2, cols=2)
    ht.alignment = WD_TABLE_ALIGNMENT.CENTER
    _set_table_borders(ht)

    _set_cell_text(ht.cell(0, 0), '제목', bold=True, size=10,
                   color=(255, 255, 255), align=WD_ALIGN_PARAGRAPH.CENTER)
    _set_cell_shading(ht.cell(0, 0), C["HDR"])
    _set_cell_text(ht.cell(0, 1), config["header_title"], bold=True, size=11,
                   color=(255, 255, 255), align=WD_ALIGN_PARAGRAPH.CENTER)
    _set_cell_shading(ht.cell(0, 1), C["HDR"])

    _set_cell_text(ht.cell(1, 0), '진료과', bold=True, size=10,
                   align=WD_ALIGN_PARAGRAPH.CENTER)
    _set_cell_shading(ht.cell(1, 0), C["SEC"])
    _set_cell_text(ht.cell(1, 1), config.get("department", "재활의학과"),
                   bold=True, size=10)

    ht.cell(0, 0).width = Cm(3)
    ht.cell(0, 1).width = Cm(15)

    doc.add_paragraph()

    # ── Main protocol table ──
    # Flatten sections into data list
    data = []
    for sec_idx, sec in enumerate(config["sections"]):
        for row_idx, row in enumerate(sec["rows"]):
            sub_label, sub_bg, content = row
            sec_label = sec["label"] if row_idx == 0 else None
            data.append((sec_label, sub_label, sub_bg, content))
        # Add separator between sections (not after last)
        if sec_idx < len(config["sections"]) - 1:
            data.append(("SEP", None, None, None))

    mt = doc.add_table(rows=0, cols=3)
    mt.alignment = WD_TABLE_ALIGNMENT.CENTER
    _set_table_borders(mt)

    section_ranges = []
    current_section_start = None
    current_section_label = None
    row_idx = 0

    for item in data:
        sec_label, sub_label, sub_bg, content = item

        if sec_label == 'SEP':
            r = mt.add_row()
            r.cells[0].merge(r.cells[1]).merge(r.cells[2])
            _set_cell_text(r.cells[0], '', size=2)
            _set_cell_shading(r.cells[0], 'FFFFFF')
            if current_section_start is not None:
                section_ranges.append((current_section_start, row_idx - 1, current_section_label))
                current_section_start = None
            row_idx += 1
            continue

        if sub_label == 'COSIGN':
            r = mt.add_row()
            _set_cell_shading(r.cells[0], C["SEC"])
            r.cells[1].merge(r.cells[2])
            cell = r.cells[1]
            cell.text = ''
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(4)
            run = p.add_run('★ 초안작성(임시저장) → 의사 cosign')
            run.bold = True
            run.font.size = Pt(11)
            run.font.name = '맑은 고딕'
            run.font.element.rPr.rFonts.set(qn('w:eastAsia'), '맑은 고딕')
            run.font.color.rgb = RGBColor(0xFF, 0x00, 0x00)
            _set_cell_shading(r.cells[1], C["WARN"])
            row_idx += 1
            continue

        r = mt.add_row()

        if sec_label is not None:
            if current_section_start is not None:
                section_ranges.append((current_section_start, row_idx - 1, current_section_label))
            current_section_start = row_idx
            current_section_label = sec_label
        _set_cell_shading(r.cells[0], C["SEC"])

        _set_cell_text(r.cells[1], sub_label, bold=True, size=9)
        if sub_bg:
            _set_cell_shading(r.cells[1], sub_bg)

        if content:
            _add_lines(r.cells[2], content, size=8)

        row_idx += 1

    # Close last section
    if current_section_start is not None:
        section_ranges.append((current_section_start, row_idx - 1, current_section_label))

    # Merge section label cells
    all_rows = mt.rows
    for start, end, label in section_ranges:
        if start == end:
            _set_cell_text(all_rows[start].cells[0], label, bold=True, size=10,
                           align=WD_ALIGN_PARAGRAPH.CENTER)
            _set_cell_shading(all_rows[start].cells[0], C["SEC"])
        else:
            merged = all_rows[start].cells[0]
            for i in range(start + 1, end + 1):
                merged = merged.merge(all_rows[i].cells[0])
            _set_cell_text(merged, label, bold=True, size=10,
                           align=WD_ALIGN_PARAGRAPH.CENTER)
            _set_cell_shading(merged, C["SEC"])

    # Column widths
    for row in mt.rows:
        if len(row.cells) >= 3:
            row.cells[0].width = Cm(2.5)
            row.cells[1].width = Cm(5.0)
            row.cells[2].width = Cm(10.5)
        elif len(row.cells) >= 2:
            row.cells[0].width = Cm(2.5)
            row.cells[1].width = Cm(15.5)

    # ── Notes (비고) ──
    if config.get("notes"):
        doc.add_paragraph()

        bt = doc.add_table(rows=1, cols=2)
        bt.alignment = WD_TABLE_ALIGNMENT.CENTER
        _set_table_borders(bt)
        bt.cell(0, 0).width = Cm(2.5)
        bt.cell(0, 1).width = Cm(15.5)

        _set_cell_text(bt.cell(0, 0), '비고', bold=True, size=10,
                       align=WD_ALIGN_PARAGRAPH.CENTER)
        _set_cell_shading(bt.cell(0, 0), C["NOTE"])

        cell = bt.cell(0, 1)
        cell.text = ''
        p = cell.paragraphs[0]
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(2)

        for text, bold, red in config["notes"]:
            run = p.add_run(text)
            run.font.size = Pt(8)
            run.font.name = '맑은 고딕'
            run.font.element.rPr.rFonts.set(qn('w:eastAsia'), '맑은 고딕')
            run.bold = bold
            if red:
                run.font.color.rgb = RGBColor(0xFF, 0x00, 0x00)

    # ── Save ──
    os.makedirs(os.path.dirname(output_path) or '.', exist_ok=True)
    doc.save(output_path)
    print(f'✅ Protocol generated: {output_path}')


# ══════════════════════════════════════════════════════════════
# Example: CPET protocol
# ══════════════════════════════════════════════════════════════

if __name__ == "__main__":
    C = COLORS

    example_config = {
        "title": "재활의학과 심폐운동부하검사 (CPET) 시행 프로토콜",
        "header_title": "심폐운동부하검사 (CPET) 시행 프로토콜",
        "department": "재활의학과",
        "sections": [
            {
                "label": "검사 전\n확인",
                "rows": [
                    ("검사 의뢰 확인", C["SUB"], ["처방전 확인 (심폐운동부하검사 오더)"]),
                    ("즉시 평가", C["SUB"], ["환자 금기사항 확인 (아래 금기 목록 참조)"]),
                    ("절대적 금기", C["CRIT"], [
                        "① 급성 심근경색 (2일 이내)",
                        "② 불안정 협심증",
                        "③ 조절되지 않는 부정맥 (혈역학적 불안정)",
                        "④ 증상 동반 중증 대동맥 협착",
                        "⑤ 조절되지 않는 심부전",
                        "⑥ 급성 폐색전 또는 폐경색",
                        "⑦ 급성 심근염/심낭염",
                        "⑧ 급성 대동맥 박리",
                    ]),
                    ("상대적 금기", C["WARN"], [
                        "① 좌주관상동맥 협착",
                        "② 중등도 판막 협착",
                        "③ 전해질 이상",
                        "④ 조절되지 않는 고혈압 (SBP >200 or DBP >110)",
                        "⑤ 빈맥성/서맥성 부정맥",
                        "⑥ 비후성 심근병증",
                        "⑦ 정신적/신체적 검사 불가",
                        "⑧ 고도 방실차단",
                    ]),
                    ("의사 Call 여부 판단", C["SUB"], [
                        ("절대적 금기 해당 → 검사 불가, 의사 Call", True),
                        ("상대적 금기 해당 → 의사 Call하여 시행 여부 확인", True),
                        ("금기 없음 → 프로토콜 수행", True),
                    ]),
                ],
            },
            # ... 나머지 섹션은 동일 패턴으로 추가
        ],
        "notes": [
            ("※ 「간호법」 제12조제2항 및 제14조에 따라 간호사의 진료지원업무가 신설", False, False),
            ("\n※ 진료지원업무 수행 간호사는 ", False, False),
            ("프로토콜 하 검사·약물의 처방 초안 작성만 가능", True, True),
            ("\n\n※ 진료과 set 처방, CP 활용 가능", False, False),
        ],
    }

    generate_protocol(example_config, "/workspace/research/pytool/example_output.docx")
