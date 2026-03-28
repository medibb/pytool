#!/usr/bin/env python3
"""
reporting_checklist.py — CONSORT/STROBE Reporting Guideline Validator
                        + Clinical Trial Registration Draft

Phase 6 of the clinical research workflow:
  1. CONSORT 2010 Checklist Validator (25 items for RCTs)
  2. STROBE Checklist Validator (22 items for observational studies)
  3. Clinical trial registration draft (ClinicalTrials.gov / CRIS format)

=== Usage ===

Python:
    from reporting_checklist import (
        validate_consort, validate_strobe,
        generate_registration_draft, generate_report,
        EXAMPLE_MANUSCRIPT_CONFIG
    )

    result = validate_consort(manuscript_config)
    result = validate_strobe(manuscript_config)
    draft = generate_registration_draft(study_config)

CLI:
    python3 reporting_checklist.py consort --config config.json -o report.docx
    python3 reporting_checklist.py strobe --config config.json -o report.docx
    python3 reporting_checklist.py register --config config.json -o draft.docx
    python3 reporting_checklist.py --example --run

Dependencies: python-docx
"""

import os
import sys
import json
import argparse
from datetime import datetime

try:
    from docx import Document
    from docx.shared import Pt, RGBColor, Inches
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    HAS_DOCX = True
except ImportError:
    HAS_DOCX = False


# ============================================================
# CONSORT 2010 Checklist (25 items)
# ============================================================

CONSORT_CHECKLIST = [
    # Title & Abstract
    {"id": "1a", "section": "Title and Abstract",
     "description": "Identification as a randomised trial in the title",
     "keywords": ["randomised", "randomized", "rct", "trial"],
     "manuscript_section": "title"},
    {"id": "1b", "section": "Title and Abstract",
     "description": "Structured summary of trial design, methods, results, conclusions (CONSORT abstract)",
     "keywords": ["abstract", "summary", "methods", "results", "conclusions"],
     "manuscript_section": "abstract"},

    # Introduction
    {"id": "2a", "section": "Introduction",
     "description": "Scientific background and explanation of rationale",
     "keywords": ["background", "rationale", "evidence"],
     "manuscript_section": "introduction"},
    {"id": "2b", "section": "Introduction",
     "description": "Specific objectives or hypotheses",
     "keywords": ["objective", "aim", "hypothesis"],
     "manuscript_section": "introduction"},

    # Methods
    {"id": "3a", "section": "Methods",
     "description": "Description of trial design including allocation ratio",
     "keywords": ["parallel", "crossover", "allocation ratio", "superiority", "non-inferiority"],
     "manuscript_section": "methods"},
    {"id": "3b", "section": "Methods",
     "description": "Important changes to methods after trial commencement with reasons",
     "keywords": ["protocol amendment", "change", "modification"],
     "manuscript_section": "methods"},
    {"id": "4a", "section": "Methods",
     "description": "Eligibility criteria for participants",
     "keywords": ["inclusion", "exclusion", "eligibility", "criteria"],
     "manuscript_section": "methods"},
    {"id": "4b", "section": "Methods",
     "description": "Settings and locations where data were collected",
     "keywords": ["setting", "hospital", "center", "location"],
     "manuscript_section": "methods"},
    {"id": "5", "section": "Methods",
     "description": "Interventions for each group with sufficient details for replication",
     "keywords": ["intervention", "procedure", "treatment", "comparator"],
     "manuscript_section": "methods"},
    {"id": "6a", "section": "Methods",
     "description": "Completely defined pre-specified primary and secondary outcomes with assessment timing",
     "keywords": ["primary outcome", "secondary outcome", "endpoint", "time point"],
     "manuscript_section": "methods"},
    {"id": "6b", "section": "Methods",
     "description": "Any changes to trial outcomes after the trial commenced, with reasons",
     "keywords": ["outcome change", "amendment"],
     "manuscript_section": "methods"},
    {"id": "7a", "section": "Methods",
     "description": "How sample size was determined",
     "keywords": ["sample size", "power", "calculation"],
     "manuscript_section": "methods"},
    {"id": "7b", "section": "Methods",
     "description": "When applicable, explanation of any interim analyses and stopping guidelines",
     "keywords": ["interim", "stopping", "futility"],
     "manuscript_section": "methods"},
    {"id": "8a", "section": "Methods",
     "description": "Method used to generate the random allocation sequence",
     "keywords": ["randomisation", "randomization", "sequence generation", "computer-generated"],
     "manuscript_section": "methods"},
    {"id": "8b", "section": "Methods",
     "description": "Type of randomisation; details of any restriction",
     "keywords": ["block", "stratified", "simple", "restriction"],
     "manuscript_section": "methods"},
    {"id": "9", "section": "Methods",
     "description": "Mechanism used to implement the random allocation sequence, concealment",
     "keywords": ["concealment", "sealed", "opaque", "central"],
     "manuscript_section": "methods"},
    {"id": "10", "section": "Methods",
     "description": "Who generated the sequence, enrolled participants, assigned to interventions",
     "keywords": ["enrolment", "assignment", "generated"],
     "manuscript_section": "methods"},
    {"id": "11a", "section": "Methods",
     "description": "If done, who was blinded after assignment and how",
     "keywords": ["blinding", "masking", "double-blind", "single-blind"],
     "manuscript_section": "methods"},
    {"id": "11b", "section": "Methods",
     "description": "If relevant, description of the similarity of interventions",
     "keywords": ["placebo", "sham", "identical", "indistinguishable"],
     "manuscript_section": "methods"},
    {"id": "12a", "section": "Methods",
     "description": "Statistical methods used to compare groups for primary/secondary outcomes",
     "keywords": ["t-test", "ANCOVA", "chi-square", "regression", "mixed model"],
     "manuscript_section": "methods"},
    {"id": "12b", "section": "Methods",
     "description": "Methods for additional analyses (e.g., subgroup, adjusted analyses)",
     "keywords": ["subgroup", "adjusted", "sensitivity", "per-protocol"],
     "manuscript_section": "methods"},

    # Results
    {"id": "13a", "section": "Results",
     "description": "For each group, numbers of participants randomly assigned, receiving treatment, completing follow-up, analysed",
     "keywords": ["flow", "randomised", "analysed", "completed", "lost to follow-up"],
     "manuscript_section": "results"},
    {"id": "13b", "section": "Results",
     "description": "For each group, losses and exclusions after randomisation, with reasons",
     "keywords": ["withdrawal", "dropout", "excluded", "discontinued"],
     "manuscript_section": "results"},
    {"id": "14a", "section": "Results",
     "description": "Dates defining recruitment and follow-up periods",
     "keywords": ["recruitment", "enrolment", "follow-up period"],
     "manuscript_section": "results"},
    {"id": "14b", "section": "Results",
     "description": "Why the trial ended or was stopped",
     "keywords": ["trial ended", "stopped", "terminated"],
     "manuscript_section": "results"},
    {"id": "15", "section": "Results",
     "description": "Baseline demographic and clinical characteristics table for each group",
     "keywords": ["table 1", "baseline", "demographic", "characteristics"],
     "manuscript_section": "results"},
    {"id": "16", "section": "Results",
     "description": "Number of participants (denominator) in each analysis group",
     "keywords": ["ITT", "intention-to-treat", "per-protocol", "denominator"],
     "manuscript_section": "results"},
    {"id": "17a", "section": "Results",
     "description": "For each primary/secondary outcome, results for each group with estimated effect size and its precision",
     "keywords": ["effect size", "confidence interval", "mean difference", "risk ratio"],
     "manuscript_section": "results"},
    {"id": "17b", "section": "Results",
     "description": "For binary outcomes, presentation of both absolute and relative effect sizes recommended",
     "keywords": ["absolute", "relative", "NNT", "risk difference"],
     "manuscript_section": "results"},
    {"id": "18", "section": "Results",
     "description": "Results of any other analyses (subgroup, adjusted, sensitivity)",
     "keywords": ["subgroup", "adjusted", "sensitivity analysis"],
     "manuscript_section": "results"},
    {"id": "19", "section": "Results",
     "description": "All important harms or unintended effects in each group",
     "keywords": ["adverse event", "harm", "safety", "side effect"],
     "manuscript_section": "results"},

    # Discussion
    {"id": "20", "section": "Discussion",
     "description": "Trial limitations, addressing sources of potential bias, imprecision",
     "keywords": ["limitation", "bias", "imprecision", "generalisability"],
     "manuscript_section": "discussion"},
    {"id": "21", "section": "Discussion",
     "description": "Generalisability (external validity) of the trial findings",
     "keywords": ["generalisability", "external validity", "applicability"],
     "manuscript_section": "discussion"},
    {"id": "22", "section": "Discussion",
     "description": "Interpretation consistent with results, balancing benefits and harms, considering other evidence",
     "keywords": ["interpretation", "clinical significance", "benefit", "harm"],
     "manuscript_section": "discussion"},

    # Other
    {"id": "23", "section": "Other",
     "description": "Registration number and name of trial registry",
     "keywords": ["registration", "NCT", "CRIS", "trial registry"],
     "manuscript_section": "other"},
    {"id": "24", "section": "Other",
     "description": "Where the full trial protocol can be accessed",
     "keywords": ["protocol", "available", "supplement", "appendix"],
     "manuscript_section": "other"},
    {"id": "25", "section": "Other",
     "description": "Sources of funding and other support; role of funders",
     "keywords": ["funding", "support", "grant", "sponsor"],
     "manuscript_section": "other"},
]


# ============================================================
# STROBE Checklist (22 items for observational studies)
# ============================================================

STROBE_CHECKLIST = [
    {"id": "1", "section": "Title and Abstract",
     "description": "Indicate the study's design with a commonly used term in the title or abstract; provide an informative abstract",
     "keywords": ["cohort", "case-control", "cross-sectional", "observational"],
     "manuscript_section": "title"},
    {"id": "2", "section": "Introduction",
     "description": "Explain the scientific background and rationale",
     "keywords": ["background", "rationale"],
     "manuscript_section": "introduction"},
    {"id": "3", "section": "Introduction",
     "description": "State specific objectives, including any prespecified hypotheses",
     "keywords": ["objective", "hypothesis"],
     "manuscript_section": "introduction"},
    {"id": "4", "section": "Methods",
     "description": "Present key elements of study design early in the paper",
     "keywords": ["study design", "prospective", "retrospective"],
     "manuscript_section": "methods"},
    {"id": "5", "section": "Methods",
     "description": "Describe the setting, locations, and relevant dates",
     "keywords": ["setting", "location", "dates", "period"],
     "manuscript_section": "methods"},
    {"id": "6", "section": "Methods",
     "description": "Give the eligibility criteria, sources and methods of selection of participants",
     "keywords": ["eligibility", "inclusion", "exclusion", "selection"],
     "manuscript_section": "methods"},
    {"id": "7", "section": "Methods",
     "description": "Clearly define all outcomes, exposures, predictors, potential confounders, and effect modifiers",
     "keywords": ["outcome", "exposure", "confounder", "variable"],
     "manuscript_section": "methods"},
    {"id": "8", "section": "Methods",
     "description": "For each variable, give sources of data and details of methods of assessment/measurement",
     "keywords": ["measurement", "assessment", "instrument", "validity"],
     "manuscript_section": "methods"},
    {"id": "9", "section": "Methods",
     "description": "Describe any efforts to address potential sources of bias",
     "keywords": ["bias", "confounding", "matching"],
     "manuscript_section": "methods"},
    {"id": "10", "section": "Methods",
     "description": "Explain how the study size was arrived at",
     "keywords": ["sample size", "power"],
     "manuscript_section": "methods"},
    {"id": "11", "section": "Methods",
     "description": "Explain how quantitative variables were handled in the analyses",
     "keywords": ["continuous", "categorical", "grouping", "cutpoint"],
     "manuscript_section": "methods"},
    {"id": "12", "section": "Methods",
     "description": "Describe all statistical methods, including those to control for confounding",
     "keywords": ["statistical", "regression", "adjustment", "model"],
     "manuscript_section": "methods"},
    {"id": "13", "section": "Results",
     "description": "Report numbers of individuals at each stage of study",
     "keywords": ["flow diagram", "eligible", "included", "excluded"],
     "manuscript_section": "results"},
    {"id": "14", "section": "Results",
     "description": "Give characteristics of study participants and information on exposures/confounders",
     "keywords": ["table 1", "characteristics", "demographic"],
     "manuscript_section": "results"},
    {"id": "15", "section": "Results",
     "description": "Indicate number of participants with missing data for each variable",
     "keywords": ["missing", "complete case", "available"],
     "manuscript_section": "results"},
    {"id": "16", "section": "Results",
     "description": "Report numbers of outcome events or summary measures",
     "keywords": ["outcome", "event", "incidence", "prevalence"],
     "manuscript_section": "results"},
    {"id": "17", "section": "Results",
     "description": "Give unadjusted estimates and, if applicable, confounder-adjusted estimates with precision",
     "keywords": ["odds ratio", "risk ratio", "confidence interval", "adjusted"],
     "manuscript_section": "results"},
    {"id": "18", "section": "Results",
     "description": "Report other analyses done (e.g., sensitivity or subgroup analyses)",
     "keywords": ["subgroup", "sensitivity", "stratified"],
     "manuscript_section": "results"},
    {"id": "19", "section": "Discussion",
     "description": "Summarise key results with reference to study objectives",
     "keywords": ["summary", "key finding", "main result"],
     "manuscript_section": "discussion"},
    {"id": "20", "section": "Discussion",
     "description": "Discuss limitations, including sources of potential bias or imprecision",
     "keywords": ["limitation", "bias", "imprecision"],
     "manuscript_section": "discussion"},
    {"id": "21", "section": "Discussion",
     "description": "Give a cautious overall interpretation of results",
     "keywords": ["interpretation", "clinical implication", "generalisability"],
     "manuscript_section": "discussion"},
    {"id": "22", "section": "Other",
     "description": "Give the source of funding and the role of the funders",
     "keywords": ["funding", "grant", "sponsor"],
     "manuscript_section": "other"},
]


# ============================================================
# Validation Engine
# ============================================================

def _validate_checklist(checklist, manuscript_config, checklist_name):
    """Generic validation engine for reporting checklists.

    Args:
        checklist: list of checklist items
        manuscript_config: dict with manuscript section content
            keys: title, abstract, introduction, methods, results, discussion, other
            values: text content or list of keywords/phrases present
        checklist_name: "CONSORT" or "STROBE"

    Returns:
        dict: {items, score, missing, addressed}
    """
    results = []
    sections = manuscript_config.get("sections", {})
    reported_keywords = manuscript_config.get("keywords_present", [])
    all_keywords_lower = [k.lower() for k in reported_keywords]

    for item in checklist:
        status = "not_addressed"
        evidence = []

        # Check if any item keywords found in manuscript sections or keyword list
        ms = item.get("manuscript_section", "")
        section_text = str(sections.get(ms, "")).lower()

        for kw in item["keywords"]:
            kw_lower = kw.lower()
            if kw_lower in section_text:
                evidence.append(f"Found '{kw}' in {ms}")
                status = "addressed"
            elif kw_lower in all_keywords_lower:
                evidence.append(f"Keyword '{kw}' reported")
                status = "addressed"

        # Check explicit item status overrides
        explicit = manuscript_config.get("checklist_status", {})
        if item["id"] in explicit:
            override = explicit[item["id"]]
            if override.get("status") == "addressed":
                status = "addressed"
                evidence.append(override.get("page", "User-marked"))
            elif override.get("status") == "not_applicable":
                status = "not_applicable"
                evidence.append("Not applicable for this study")

        results.append({
            "id": item["id"],
            "section": item["section"],
            "description": item["description"],
            "status": status,
            "evidence": evidence,
        })

    addressed = [r for r in results if r["status"] == "addressed"]
    missing = [r for r in results if r["status"] == "not_addressed"]
    na = [r for r in results if r["status"] == "not_applicable"]
    checkable = len(results) - len(na)

    return {
        "checklist_name": checklist_name,
        "items": results,
        "score": {
            "addressed": len(addressed),
            "missing": len(missing),
            "not_applicable": len(na),
            "total": len(results),
            "checkable": checkable,
            "percentage": round(len(addressed) / max(checkable, 1) * 100, 1),
        },
        "addressed": addressed,
        "missing": missing,
    }


def validate_consort(manuscript_config):
    """Validate manuscript against CONSORT 2010 checklist."""
    return _validate_checklist(CONSORT_CHECKLIST, manuscript_config, "CONSORT 2010")


def validate_strobe(manuscript_config):
    """Validate manuscript against STROBE checklist."""
    return _validate_checklist(STROBE_CHECKLIST, manuscript_config, "STROBE")


# ============================================================
# Clinical Trial Registration Draft
# ============================================================

def generate_registration_draft(study_config, output_path="trial_registration_draft.docx"):
    """Generate clinical trial registration draft.

    Follows ClinicalTrials.gov / CRIS registration fields.
    Auto-fills from study config (scaffold + SAP).
    """
    if not HAS_DOCX:
        raise ImportError("python-docx required.")

    doc = Document()
    study = study_config.get("study", {})

    doc.add_heading("Clinical Trial Registration Draft", level=0)
    doc.add_paragraph(f"Generated: {datetime.now().strftime('%Y-%m-%d')}")
    doc.add_paragraph("Format: ClinicalTrials.gov / CRIS compatible fields")
    doc.add_paragraph("")

    # Registration fields
    fields = _build_registration_fields(study_config)

    table = doc.add_table(rows=len(fields) + 1, cols=3)
    table.style = "Table Grid"

    # Header
    for i, h in enumerate(["Field", "Value", "Source"]):
        run = table.rows[0].cells[i].paragraphs[0].add_run(h)
        run.bold = True
        run.font.size = Pt(10)

    for idx, (field_name, value, source) in enumerate(fields, 1):
        row = table.rows[idx]
        row.cells[0].paragraphs[0].add_run(field_name).font.size = Pt(9)

        run = row.cells[1].paragraphs[0].add_run(str(value))
        run.font.size = Pt(9)
        if not value or value == "—":
            run.font.color.rgb = RGBColor(0xCC, 0x00, 0x00)

        row.cells[2].paragraphs[0].add_run(source).font.size = Pt(8)

    doc.save(output_path)
    return output_path


def _build_registration_fields(config):
    """Build registration field list from config."""
    study = config.get("study", {})
    design = config.get("design", {})
    po = config.get("primary_outcome", {})
    arms = design.get("arms", [])
    inclusion = config.get("inclusion", [])
    exclusion = config.get("exclusion", [])
    secondary = config.get("secondary_outcomes", [])
    followup = config.get("followup", {})

    # Study type mapping
    design_type = design.get("type", "")
    if "rct" in design_type:
        study_type = "Interventional"
        phase_text = "Not Applicable (device/procedure study)"
    else:
        study_type = "Observational"
        phase_text = "Not Applicable"

    # Blinding
    blinding_map = {
        "open_label": "None (Open Label)",
        "single_blind_assessor": "Single (Outcomes Assessor)",
        "double_blind": "Double (Participant, Outcomes Assessor)",
        "triple_blind": "Triple (Participant, Care Provider, Outcomes Assessor)",
    }
    blinding = blinding_map.get(design.get("blinding", ""), design.get("blinding", "—"))

    # Arms
    arm_text = "\n".join(f"  {i+1}. {a.get('name', '')}" for i, a in enumerate(arms)) if arms else "—"

    # Sample size
    try:
        from scipy.stats import norm
        import math
        mcid = po.get("mcid", 0)
        sd = po.get("sd", 0)
        alpha = po.get("alpha", 0.05)
        power = po.get("power", 0.80)
        dropout = po.get("dropout_rate", 0)
        if mcid and sd:
            z_a = norm.ppf(1 - alpha/2)
            z_b = norm.ppf(power)
            n = math.ceil(2 * ((z_a + z_b) * sd / mcid) ** 2)
            n_adj = math.ceil(n / (1 - dropout)) if dropout else n
            total = n_adj * len(arms) if arms else n_adj * 2
            sample_text = str(total)
        else:
            sample_text = "—"
    except Exception:
        sample_text = "—"

    fields = [
        # Descriptive Information
        ("Brief Title", study.get("title_en", study.get("title", "—")), "study.title"),
        ("Official Title", study.get("title", study.get("title_ko", "—")), "study.title"),
        ("Brief Summary", config.get("objectives", {}).get("primary", "—"), "objectives"),
        ("Detailed Description", config.get("background", "—"), "background"),

        # Study Design
        ("Study Type", study_type, "design.type"),
        ("Study Phase", phase_text, "design.type"),
        ("Study Design", f"Allocation: Randomized\n  Intervention Model: Parallel Assignment\n  Masking: {blinding}", "design"),
        ("Number of Arms", str(len(arms)) if arms else "—", "design.arms"),
        ("Arms", arm_text, "design.arms"),

        # Eligibility
        ("Eligibility Criteria — Inclusion",
         "\n".join(f"  - {c}" for c in inclusion) if inclusion else "—",
         "inclusion"),
        ("Eligibility Criteria — Exclusion",
         "\n".join(f"  - {c}" for c in exclusion) if exclusion else "—",
         "exclusion"),
        ("Sex/Gender", "All", "—"),
        ("Age Limits", _extract_age_limits(inclusion), "inclusion"),

        # Outcomes
        ("Primary Outcome Measure", po.get("name", "—"), "primary_outcome"),
        ("Primary Outcome Time Frame",
         followup.get("visits", ["—"])[-1] if followup.get("visits") else "—",
         "followup"),
        ("Secondary Outcome Measures",
         "\n".join(f"  {i+1}. {s.get('name', s) if isinstance(s, dict) else s}"
                   for i, s in enumerate(secondary)) if secondary else "—",
         "secondary_outcomes"),

        # Enrollment
        ("Estimated Enrollment", sample_text, "primary_outcome (calculated)"),

        # Contacts & Locations
        ("Responsible Party", study.get("pi", "—"), "study.pi"),
        ("Sponsor", study.get("department", study.get("sponsor", "—")), "study.department"),
        ("Study Location", study.get("department", "—"), "study.department"),
        ("Country", "Korea, Republic of", "—"),

        # Identifiers
        ("Protocol ID", study.get("protocol_number", "—"), "study.protocol_number"),
    ]

    return fields


def _extract_age_limits(inclusion):
    """Try to extract age range from inclusion criteria."""
    import re
    for criterion in inclusion:
        match = re.search(r'(\d+)[–\-~](\d+)\s*(?:세|years|yr)', criterion)
        if match:
            return f"{match.group(1)} Years to {match.group(2)} Years"
        match = re.search(r'≥?\s*(\d+)\s*(?:세|years|yr)', criterion)
        if match:
            return f"{match.group(1)} Years and older"
    return "18 Years and older"


# ============================================================
# Report Generator (docx)
# ============================================================

def generate_report(manuscript_config, study_config=None, output_path="reporting_checklist_report.docx"):
    """Generate comprehensive reporting checklist report.

    Includes CONSORT or STROBE validation + trial registration draft.
    """
    if not HAS_DOCX:
        raise ImportError("python-docx required.")

    doc = Document()
    doc.add_heading("Reporting Guideline Compliance Report", level=0)
    doc.add_paragraph(f"Generated: {datetime.now().strftime('%Y-%m-%d')}")
    doc.add_paragraph(f"Tool: reporting_checklist.py (pytool Phase 6)")
    doc.add_paragraph("")

    # Determine which checklist to use
    design_type = manuscript_config.get("design_type", "rct")
    if "rct" in design_type.lower() or "randomized" in design_type.lower():
        result = validate_consort(manuscript_config)
    else:
        result = validate_strobe(manuscript_config)

    score = result["score"]

    # Summary
    doc.add_heading(f"1. {result['checklist_name']} Checklist Validation", level=1)
    p = doc.add_paragraph()
    p.add_run("Score: ").bold = True
    p.add_run(f"{score['addressed']}/{score['checkable']} items addressed "
              f"({score['percentage']}%)")

    # Checklist table
    table = doc.add_table(rows=len(result["items"]) + 1, cols=4)
    table.style = "Table Grid"
    for i, h in enumerate(["Item", "Section", "Description", "Status"]):
        run = table.rows[0].cells[i].paragraphs[0].add_run(h)
        run.bold = True
        run.font.size = Pt(9)

    status_map = {"addressed": "✅", "not_addressed": "❌", "not_applicable": "N/A"}
    for idx, item in enumerate(result["items"], 1):
        row = table.rows[idx]
        for j, text in enumerate([
            item["id"], item["section"],
            item["description"][:80],
            status_map.get(item["status"], "?")
        ]):
            run = row.cells[j].paragraphs[0].add_run(str(text))
            run.font.size = Pt(8)
            if item["status"] == "not_addressed":
                run.font.color.rgb = RGBColor(0xCC, 0x00, 0x00)

    # Missing items
    if result["missing"]:
        doc.add_heading("1.1 Action Required: Missing Items", level=2)
        for m in result["missing"]:
            p = doc.add_paragraph(style="List Bullet")
            run = p.add_run(f"Item {m['id']}: ")
            run.bold = True
            p.add_run(m["description"])

    # Trial registration
    if study_config:
        doc.add_page_break()
        doc.add_heading("2. Clinical Trial Registration Draft", level=1)
        fields = _build_registration_fields(study_config)
        reg_table = doc.add_table(rows=len(fields) + 1, cols=2)
        reg_table.style = "Table Grid"
        reg_table.rows[0].cells[0].paragraphs[0].add_run("Field").bold = True
        reg_table.rows[0].cells[1].paragraphs[0].add_run("Value").bold = True
        for idx, (field, value, _) in enumerate(fields, 1):
            row = reg_table.rows[idx]
            row.cells[0].paragraphs[0].add_run(field).font.size = Pt(9)
            run = row.cells[1].paragraphs[0].add_run(str(value))
            run.font.size = Pt(9)

    doc.save(output_path)
    return {
        "output_path": output_path,
        "checklist": result["checklist_name"],
        "score": score,
    }


# ============================================================
# Example Config
# ============================================================

EXAMPLE_MANUSCRIPT_CONFIG = {
    "design_type": "rct",
    "sections": {
        "title": "A Randomized Controlled Trial Comparing Anterior vs Posterior Approach "
                 "Ultrasound-guided Hydrodilatation for Frozen Shoulder",
        "abstract": "Background: Hydrodilatation is effective for frozen shoulder. Methods: "
                    "We randomized 98 patients to anterior or posterior approach. Results: "
                    "Both groups showed significant improvement in SPADI. Conclusions: "
                    "Anterior approach was non-inferior to posterior approach.",
        "introduction": "Adhesive capsulitis affects 2-5% of the general population. "
                        "The scientific background shows hydrodilatation improves outcomes. "
                        "Our objective was to compare anterior vs posterior approach.",
        "methods": "This was a parallel-group, single-blind (assessor-blind) randomized "
                   "controlled trial. Inclusion criteria: age 40-75, ROM limitation >=30 degrees. "
                   "Exclusion criteria: secondary causes, recent steroid injection. "
                   "The primary outcome was SPADI total score change at 3 months. "
                   "Sample size was calculated using alpha=0.05, power=0.80, MCID=13, SD=20. "
                   "Block randomization with stratification by sex was performed. "
                   "Allocation concealment used sealed opaque envelopes. "
                   "Independent samples t-test was used for primary analysis. "
                   "Subgroup analyses by sex and diabetes status were pre-specified. "
                   "Per-protocol sensitivity analysis was also performed.",
        "results": "Figure 1 shows the flow diagram. 98 participants were randomised. "
                   "Table 1 presents baseline characteristics. "
                   "The mean SPADI change was -25.3 (95% CI: -30.1 to -20.5) in the anterior "
                   "group vs -23.8 (95% CI: -28.7 to -18.9) in the posterior group. "
                   "The mean difference was -1.5 (95% CI: -8.2 to 5.2, p=0.66). "
                   "No serious adverse events occurred. One vasovagal reaction in each group.",
        "discussion": "Our findings suggest both approaches are similarly effective. "
                      "Limitations include single-center design and relatively small sample. "
                      "The results may be generalisable to similar clinical settings. "
                      "These findings are consistent with prior evidence on hydrodilatation.",
        "other": "Trial registration: CRIS KCT0009XXX. "
                 "The full protocol is available as a supplementary appendix. "
                 "This study was funded by departmental research funds.",
    },
    "keywords_present": [
        "randomized", "parallel", "allocation ratio", "SPADI",
        "inclusion", "exclusion", "sample size", "power",
        "block randomization", "sealed opaque envelopes",
        "single-blind", "t-test", "ANCOVA", "subgroup",
        "flow diagram", "Table 1", "confidence interval",
        "adverse event", "limitation", "generalisability",
        "CRIS", "protocol", "funding",
    ],
}

EXAMPLE_STUDY_CONFIG = {
    "study": {
        "title": "Comparison of Anterior vs Posterior Approach Hydrodilatation for Frozen Shoulder",
        "title_en": "Anterior vs Posterior Hydrodilatation for Adhesive Capsulitis: An RCT",
        "title_ko": "동결견 전방 vs 후방 접근 수압팽창술 비교",
        "pi": "Jaehyun Lee, MD",
        "department": "Dept. of Rehabilitation Medicine, Pusan National University Hospital",
        "protocol_number": "PNUH-2026-IRB-XXX",
    },
    "design": {
        "type": "rct_parallel",
        "arms": [
            {"name": "Anterior approach hydrodilatation"},
            {"name": "Posterior approach hydrodilatation"},
        ],
        "blinding": "single_blind_assessor",
    },
    "primary_outcome": {
        "name": "SPADI total score change at 3 months",
        "mcid": 13, "sd": 20, "alpha": 0.05, "power": 0.80,
        "dropout_rate": 0.20,
    },
    "secondary_outcomes": [
        {"name": "Shoulder ROM change (ER, IR, ABD, Flexion)"},
        {"name": "Pain NRS change (rest, activity, night)"},
        {"name": "PGA change"},
    ],
    "inclusion": [
        "40-75세 성인", "Passive ROM 제한 ≥ 30° (2방향 이상)",
        "Primary adhesive capsulitis", "증상 기간 ≥ 4주",
    ],
    "exclusion": [
        "이차성 원인", "최근 3개월 이내 스테로이드 주사",
        "조절되지 않는 당뇨 (HbA1c > 9%)",
    ],
    "followup": {
        "visits": ["Baseline", "30min", "2wk", "4wk", "8wk", "12wk"],
    },
    "background": "Adhesive capsulitis affects 2-5% of the general population...",
    "objectives": {"primary": "Compare anterior vs posterior hydrodilatation effectiveness"},
}


# ============================================================
# CLI
# ============================================================

def main():
    parser = argparse.ArgumentParser(
        description="CONSORT/STROBE Checklist Validator + Trial Registration Draft"
    )
    subparsers = parser.add_subparsers(dest="command")

    # consort
    c = subparsers.add_parser("consort", help="Validate against CONSORT 2010")
    c.add_argument("--config", "-c", help="Manuscript config JSON")
    c.add_argument("--output", "-o", default="consort_report.docx")

    # strobe
    s = subparsers.add_parser("strobe", help="Validate against STROBE")
    s.add_argument("--config", "-c", help="Manuscript config JSON")
    s.add_argument("--output", "-o", default="strobe_report.docx")

    # register
    r = subparsers.add_parser("register", help="Generate trial registration draft")
    r.add_argument("--config", "-c", help="Study config JSON")
    r.add_argument("--output", "-o", default="trial_registration_draft.docx")

    # example
    parser.add_argument("--example", action="store_true", help="Use example configs")
    parser.add_argument("--run", action="store_true", help="Run with examples")

    args = parser.parse_args()

    if args.example and args.run:
        # Run all three with examples
        print("=" * 60)
        print("  CONSORT 2010 Validation")
        print("=" * 60)
        result = validate_consort(EXAMPLE_MANUSCRIPT_CONFIG)
        sc = result["score"]
        print(f"  Score: {sc['addressed']}/{sc['checkable']} ({sc['percentage']}%)")
        if result["missing"]:
            print(f"  Missing items ({sc['missing']}):")
            for m in result["missing"]:
                print(f"    ❌ {m['id']}: {m['description'][:60]}")

        print()
        report = generate_report(
            EXAMPLE_MANUSCRIPT_CONFIG,
            EXAMPLE_STUDY_CONFIG,
            "reporting_checklist_report.docx"
        )
        print(f"📄 Full report: {report['output_path']}")

        reg_path = generate_registration_draft(
            EXAMPLE_STUDY_CONFIG,
            "trial_registration_draft.docx"
        )
        print(f"📋 Registration draft: {reg_path}")
        return

    if args.example and not args.run:
        print(json.dumps({
            "manuscript_config": EXAMPLE_MANUSCRIPT_CONFIG,
            "study_config": EXAMPLE_STUDY_CONFIG,
        }, indent=2, ensure_ascii=False))
        return

    if args.command == "consort":
        cfg = _load_config(args.config) if args.config else EXAMPLE_MANUSCRIPT_CONFIG
        result = validate_consort(cfg)
        _print_result(result)
    elif args.command == "strobe":
        cfg = _load_config(args.config) if args.config else EXAMPLE_MANUSCRIPT_CONFIG
        result = validate_strobe(cfg)
        _print_result(result)
    elif args.command == "register":
        cfg = _load_config(args.config) if args.config else EXAMPLE_STUDY_CONFIG
        path = generate_registration_draft(cfg, args.output)
        print(f"📋 Registration draft saved: {path}")
    else:
        parser.print_help()
        print("\nUse --example --run to test all features.")


def _load_config(path):
    with open(path) as f:
        return json.load(f)


def _print_result(result):
    sc = result["score"]
    print(f"\n{result['checklist_name']} Validation: "
          f"{sc['addressed']}/{sc['checkable']} ({sc['percentage']}%)")
    if result["missing"]:
        print(f"\nMissing items ({sc['missing']}):")
        for m in result["missing"]:
            print(f"  ❌ {m['id']}: {m['description']}")


if __name__ == "__main__":
    main()
