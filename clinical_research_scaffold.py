#!/usr/bin/env python3
"""
clinical_research_scaffold.py — 임상연구 문서 3종 자동 생성기

FSHD 동결견 연구에서 검증된 "Practice-first, Research-seamless" 프레임워크를
범용화한 도구. config dict 하나로 연구별 맞춤 문서 3종을 생성합니다.

=== 핵심 철학 ===
1. 예진지: 임상 프로세스에 연구 데이터 수집을 자연스럽게 통합
2. Quick Sheet: 시술/검사실에서만 얻을 수 있는 데이터만 (예진과 중복 제거)
3. CRF: 마스터 문서 — 예진 + Quick Sheet 데이터를 통합 정리

=== 워크플로우 ===
외래 예진(전공의+환자) → 시술/검사실(PI) → 사후 정리(CRF)
     예진지              Quick Sheet         CRF

=== 사용법 ===

from clinical_research_scaffold import generate_research_docs

config = {
    "study": { ... },       # 연구 기본 정보
    "subjects": { ... },    # 대상자 식별 필드
    "enrollment": [ ... ],  # 등록 시 수집 항목 (인구통계 등)
    "inclusion": [ ... ],   # 선정 기준
    "exclusion": [ ... ],   # 제외 기준
    "baseline": { ... },    # Baseline 평가 항목 (예진에서 수집)
    "procedure": { ... },   # 시술/검사 항목 (Quick Sheet에서 수집)
    "followup": { ... },    # 추적 관찰
}

generate_research_docs(config, output_dir="./output")
# → 예진지.docx, Quick_Sheet.docx, CRF.docx 생성

CLI:
    python3 clinical_research_scaffold.py config.yaml [--output-dir ./output]
    python3 clinical_research_scaffold.py --example > example_config.yaml

의존성: python-docx, pyyaml (optional, for YAML config)
"""

import os
import sys
import json
import argparse
from datetime import datetime
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from lxml import etree

# ============================================================
# XML helpers
# ============================================================
W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"


def _qn(tag):
    """Convert 'w:xxx' to full namespace URI."""
    return f"{{{W}}}{tag.split(':')[-1]}"


def _set_cell_shading(cell, color_hex):
    """Set cell background color."""
    tc = cell._tc
    tcPr = tc.find(_qn("w:tcPr"))
    if tcPr is None:
        tcPr = etree.SubElement(tc, _qn("w:tcPr"))
        tc.insert(0, tcPr)
    shd = tcPr.find(_qn("w:shd"))
    if shd is None:
        shd = etree.SubElement(tcPr, _qn("w:shd"))
    shd.set(_qn("w:fill"), color_hex)
    shd.set(_qn("w:val"), "clear")


def _set_table_borders(table):
    """Set black borders on all table edges."""
    tblPr = table._tbl.find(_qn("w:tblPr"))
    if tblPr is None:
        tblPr = etree.SubElement(table._tbl, _qn("w:tblPr"))
    borders = tblPr.find(_qn("w:tblBorders"))
    if borders is None:
        borders = etree.SubElement(tblPr, _qn("w:tblBorders"))
    for name in ["top", "left", "bottom", "right", "insideH", "insideV"]:
        el = borders.find(_qn(f"w:{name}"))
        if el is None:
            el = etree.SubElement(borders, _qn(f"w:{name}"))
        el.set(_qn("w:val"), "single")
        el.set(_qn("w:sz"), "4")
        el.set(_qn("w:space"), "0")
        el.set(_qn("w:color"), "000000")


def _set_table_full_width(table):
    """Set table width to 100% of page."""
    tblPr = table._tbl.find(_qn("w:tblPr"))
    if tblPr is None:
        tblPr = etree.SubElement(table._tbl, _qn("w:tblPr"))
    tblW = tblPr.find(_qn("w:tblW"))
    if tblW is None:
        tblW = etree.SubElement(tblPr, _qn("w:tblW"))
    tblW.set(_qn("w:w"), "5000")
    tblW.set(_qn("w:type"), "pct")


# ============================================================
# Data source color scheme
# ============================================================
COLORS = {
    "patient": "FFFF99",   # 노랑 — 환자 직접 확인
    "examiner": "CCE5FF",  # 파랑 — 검사자/전공의 측정
    "emr": "CCE5FF",       # 파랑 — EMR 후입력
    "calc": "CCFFCC",      # 초록 — 자동 계산
    "procedure": "FFCCCC", # 빨강 — 시술/검사 중 기록
    "section": "D9D9D9",   # 회색 — 섹션 헤더
}


# ============================================================
# Cell/row helpers
# ============================================================
def _set_cell(cell, text, bold=False, size=9, align=None, color=None):
    """Set cell text with formatting."""
    cell.text = ""
    p = cell.paragraphs[0]
    if align:
        p.alignment = align
    p.paragraph_format.space_before = Pt(1)
    p.paragraph_format.space_after = Pt(1)
    run = p.add_run(text)
    run.font.size = Pt(size)
    run.font.name = "Arial"
    if bold:
        run.bold = True
    if color:
        _set_cell_shading(cell, color)
    return run


def _add_section_row(table, row_idx, text, color=None):
    """Merge all cells in a row and create a section header."""
    row = table.rows[row_idx]
    ncols = len(row.cells)
    if ncols > 1:
        row.cells[0].merge(row.cells[ncols - 1])
    _set_cell(row.cells[0], text, bold=True, size=10,
              color=color or COLORS["section"])


def _add_item_row(table, row_idx, label, value_hint, source="patient"):
    """Fill a data row: label | (blank for recording) | value hint."""
    row = table.rows[row_idx]
    color = COLORS.get(source)
    _set_cell(row.cells[0], label, bold=True, size=8, color=color)
    if len(row.cells) > 2:
        _set_cell(row.cells[1], "", size=8, color=color)
        _set_cell(row.cells[2], value_hint, size=7, color=color)
    else:
        _set_cell(row.cells[1], value_hint, size=8, color=color)


def _format_table(table):
    """Apply standard formatting to a table."""
    _set_table_borders(table)
    _set_table_full_width(table)


# ============================================================
# Document generators
# ============================================================

def _build_preexam(config, output_path):
    """
    Generate 예진지 (Pre-examination form).

    환자가 직접 작성하는 항목 + 전공의가 측정하는 항목을 한 장에 배치.
    baseline 섹션의 모든 항목이 여기에 포함됨.
    """
    doc = Document()
    for sec in doc.sections:
        sec.top_margin = Cm(1.5)
        sec.bottom_margin = Cm(1.0)
        sec.left_margin = Cm(1.5)
        sec.right_margin = Cm(1.5)

    study = config["study"]

    # Title
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run(f'{study["department"]} {study["preexam_title"]}')
    run.bold = True
    run.font.size = Pt(13)

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = subtitle.add_run(study.get("subtitle", ""))
    run.font.size = Pt(9)
    run.italic = True

    # Subject identification
    subj = config.get("subjects", {})
    fields = subj.get("fields", ["이름", "등록번호", "날짜"])
    id_line = "  /  ".join(f"{f}: ________" for f in fields)
    p = doc.add_paragraph(id_line)
    p.runs[0].font.size = Pt(9) if p.runs else None

    # Color legend
    legend_items = config.get("legend", {}).get("preexam", [
        ("■ 환자 작성", "patient"),
        ("■ 검사자 측정", "examiner"),
        ("■ 자동 계산", "calc"),
    ])
    legend_p = doc.add_paragraph()
    legend_p.paragraph_format.space_after = Pt(2)
    for label, _ in legend_items:
        run = legend_p.add_run(f"{label}  ")
        run.font.size = Pt(8)

    # Build sections from baseline config
    baseline = config.get("baseline", {})
    sections = baseline.get("sections", [])

    for section in sections:
        sec_name = section["name"]
        items = section["items"]

        # Count rows: 1 header + N items
        nrows = 1 + len(items)
        ncols = section.get("columns", 3)
        table = doc.add_table(rows=nrows, cols=ncols)
        _format_table(table)

        _add_section_row(table, 0, sec_name, section.get("color"))

        for i, item in enumerate(items):
            _add_item_row(
                table, i + 1,
                item["label"],
                item.get("hint", ""),
                item.get("source", "patient"),
            )

        doc.add_paragraph()  # spacing

    # Footer
    footer = doc.add_paragraph()
    footer.paragraph_format.space_before = Pt(6)
    emr_items = config.get("emr_post_entry", [])
    calc_items = config.get("auto_calc", [])
    notes = []
    if emr_items:
        notes.append(f'EMR 후입력: {", ".join(emr_items)}')
    if calc_items:
        notes.append(f'자동계산: {", ".join(calc_items)}')
    if notes:
        run = footer.add_run("※ " + "  |  ".join(notes))
        run.font.size = Pt(7)
        run.italic = True

    # Examiner signature
    sig = doc.add_paragraph()
    sig.paragraph_format.space_before = Pt(8)
    run = sig.add_run("예진자: __________  서명: __________  날짜: ____/____/____")
    run.font.size = Pt(9)

    doc.save(output_path)
    return output_path


def _build_quicksheet(config, output_path):
    """
    Generate Quick Sheet (시술/검사실 전용).

    예진에서 수집한 항목은 제외하고, 시술/검사실에서만 얻을 수 있는 데이터만 포함.
    """
    doc = Document()
    for sec in doc.sections:
        sec.top_margin = Cm(1.5)
        sec.bottom_margin = Cm(1.0)
        sec.left_margin = Cm(1.5)
        sec.right_margin = Cm(1.5)

    study = config["study"]

    # Title
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run(study.get("quicksheet_title", f'{study["short_name"]} Quick Sheet'))
    run.bold = True
    run.font.size = Pt(13)

    # Subject ID line
    subj = config.get("subjects", {})
    fields = subj.get("fields", ["Subject ID", "날짜"])
    id_line = "  ".join(f"{f}: ________" for f in fields)
    p = doc.add_paragraph(id_line)
    if p.runs:
        p.runs[0].font.size = Pt(9)

    # Legend
    legend_items = config.get("legend", {}).get("quicksheet", [
        ("■ 시술/검사 중 기록", "procedure"),
    ])
    legend_p = doc.add_paragraph()
    legend_p.paragraph_format.space_after = Pt(2)
    for label, _ in legend_items:
        run = legend_p.add_run(f"{label}  ")
        run.font.size = Pt(8)

    # Build from procedure config
    procedure = config.get("procedure", {})
    sections = procedure.get("sections", [])

    for section in sections:
        sec_name = section["name"]
        items = section["items"]

        nrows = 1 + len(items)
        ncols = section.get("columns", 3)
        table = doc.add_table(rows=nrows, cols=ncols)
        _format_table(table)

        _add_section_row(table, 0, sec_name,
                         section.get("color", COLORS["procedure"]))

        for i, item in enumerate(items):
            _add_item_row(
                table, i + 1,
                item["label"],
                item.get("hint", ""),
                item.get("source", "procedure"),
            )

        doc.add_paragraph()

    # Footer — EMR note
    emr_items = config.get("emr_post_entry", [])
    if emr_items:
        note = doc.add_paragraph()
        run = note.add_run(f'※ EMR 후입력 항목: {", ".join(emr_items)}')
        run.font.size = Pt(7)
        run.italic = True

    doc.save(output_path)
    return output_path


def _build_crf(config, output_path):
    """
    Generate CRF (Case Report Form) — 마스터 문서.

    전체 연구 프로토콜을 담은 종합 문서. 예진 + Quick Sheet 데이터를 통합.
    """
    doc = Document()
    for sec in doc.sections:
        sec.top_margin = Cm(2.0)
        sec.bottom_margin = Cm(1.5)
        sec.left_margin = Cm(2.0)
        sec.right_margin = Cm(2.0)

    study = config["study"]

    # Title page
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_before = Pt(60)
    run = title.add_run("Case Report Form (CRF)")
    run.bold = True
    run.font.size = Pt(18)

    study_title = doc.add_paragraph()
    study_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = study_title.add_run(study.get("title_ko", study["short_name"]))
    run.font.size = Pt(12)

    if study.get("title_en"):
        en_title = doc.add_paragraph()
        en_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = en_title.add_run(study["title_en"])
        run.font.size = Pt(10)
        run.italic = True

    pi_info = doc.add_paragraph()
    pi_info.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pi_info.paragraph_format.space_before = Pt(20)
    run = pi_info.add_run(f'PI: {study.get("pi", "")} ({study.get("department", "")})')
    run.font.size = Pt(10)

    doc.add_page_break()

    # --- CRF sections ---
    section_num = 0

    # 1. Subject Consented
    section_num += 1
    doc.add_heading(f"{section_num}. Subject Consented", level=1)
    subj = config.get("subjects", {})
    fields = subj.get("crf_fields", subj.get("fields", ["Subject ID", "MRN", "Name", "Date"]))
    table = doc.add_table(rows=len(fields), cols=2)
    _format_table(table)
    for i, field in enumerate(fields):
        _set_cell(table.rows[i].cells[0], field, bold=True, size=9)
        _set_cell(table.rows[i].cells[1], "", size=9)

    # 2. Enrollment
    section_num += 1
    doc.add_heading(f"{section_num}. Initial Enrollment", level=1)
    enrollment = config.get("enrollment", [])
    if enrollment:
        table = doc.add_table(rows=len(enrollment), cols=3)
        _format_table(table)
        for i, item in enumerate(enrollment):
            src = item.get("source", "emr")
            color = COLORS.get(src)
            _set_cell(table.rows[i].cells[0], item["label"], bold=True, size=9, color=color)
            _set_cell(table.rows[i].cells[1], "", size=9, color=color)
            _set_cell(table.rows[i].cells[2], item.get("hint", ""), size=8, color=color)

    # 3. Inclusion
    section_num += 1
    inclusion = config.get("inclusion", [])
    if inclusion:
        doc.add_heading(f'{section_num}. Inclusion Criteria (모두 "예" 시 적격)', level=1)
        for item in inclusion:
            p = doc.add_paragraph(style="List Number")
            run = p.add_run(f"☐ {item}")
            run.font.size = Pt(9)

    # 4. Exclusion
    section_num += 1
    exclusion = config.get("exclusion", [])
    if exclusion:
        doc.add_heading(f'{section_num}. Exclusion Criteria (하나라도 "예" 시 제외)', level=1)
        for item in exclusion:
            p = doc.add_paragraph(style="List Number")
            run = p.add_run(f"☐ {item}")
            run.font.size = Pt(9)

    # 5. Baseline Assessment (예진에서 수집)
    section_num += 1
    doc.add_heading(f"{section_num}. Baseline Assessment", level=1)
    p = doc.add_paragraph()
    run = p.add_run("※ 외래 예진에서 수집 → 예진지 참조")
    run.font.size = Pt(9)
    run.italic = True

    baseline = config.get("baseline", {})
    for si, section in enumerate(baseline.get("sections", []), 1):
        doc.add_heading(f"{section_num}.{si} {section['name']}", level=2)
        items = section["items"]
        table = doc.add_table(rows=len(items), cols=3)
        _format_table(table)
        for i, item in enumerate(items):
            src = item.get("source", "patient")
            color = COLORS.get(src)
            _set_cell(table.rows[i].cells[0], item["label"], bold=True, size=9, color=color)
            _set_cell(table.rows[i].cells[1], "", size=9, color=color)
            _set_cell(table.rows[i].cells[2], item.get("hint", ""), size=8, color=color)

    # 6+ Procedure sections (Quick Sheet에서 수집)
    procedure = config.get("procedure", {})
    for si, section in enumerate(procedure.get("sections", []), 1):
        section_num += 1
        doc.add_heading(f"{section_num}. {section['name']}", level=1)
        if si == 1:
            p = doc.add_paragraph()
            run = p.add_run("※ 시술/검사실에서 수집 → Quick Sheet 참조")
            run.font.size = Pt(9)
            run.italic = True

        items = section["items"]
        if items:
            table = doc.add_table(rows=len(items), cols=3)
            _format_table(table)
            for i, item in enumerate(items):
                src = item.get("source", "procedure")
                color = COLORS.get(src)
                _set_cell(table.rows[i].cells[0], item["label"], bold=True, size=9, color=color)
                _set_cell(table.rows[i].cells[1], "", size=9, color=color)
                _set_cell(table.rows[i].cells[2], item.get("hint", ""), size=8, color=color)

        # SOP paragraphs
        for sop in section.get("sop_text", []):
            p = doc.add_paragraph()
            run = p.add_run(sop)
            run.font.size = Pt(9)

    # Follow-up
    followup = config.get("followup", {})
    if followup:
        section_num += 1
        doc.add_heading(f"{section_num}. Follow-up Schedule", level=1)
        visits = followup.get("visits", [])
        measures = followup.get("measures", [])
        if visits:
            table = doc.add_table(rows=len(measures) + 1, cols=len(visits) + 1)
            _format_table(table)
            _set_cell(table.rows[0].cells[0], "항목", bold=True, size=8)
            for j, visit in enumerate(visits):
                _set_cell(table.rows[0].cells[j + 1], visit, bold=True, size=7,
                          align=WD_ALIGN_PARAGRAPH.CENTER)
            for i, measure in enumerate(measures):
                _set_cell(table.rows[i + 1].cells[0], measure, size=8)
                for j in range(len(visits)):
                    _set_cell(table.rows[i + 1].cells[j + 1], "○", size=8,
                              align=WD_ALIGN_PARAGRAPH.CENTER)

    # Adverse Events
    section_num += 1
    doc.add_heading(f"{section_num}. Adverse Events", level=1)
    ae_cols = ["Event", "Date", "Severity", "Causality", "Action", "Outcome"]
    table = doc.add_table(rows=2, cols=len(ae_cols))
    _format_table(table)
    for j, col in enumerate(ae_cols):
        _set_cell(table.rows[0].cells[j], col, bold=True, size=8,
                  color=COLORS["section"])
        _set_cell(table.rows[1].cells[j], "", size=8)

    # Study Completion
    section_num += 1
    doc.add_heading(f"{section_num}. Study Completion / Withdrawal", level=1)
    comp_rows = [
        ("Study completion", "☐ Completed  ☐ Withdrawn  ☐ Lost to F/U"),
        ("Date", "____/____/____"),
        ("Withdrawal reason", ""),
    ]
    table = doc.add_table(rows=len(comp_rows), cols=2)
    _format_table(table)
    for i, (label, hint) in enumerate(comp_rows):
        _set_cell(table.rows[i].cells[0], label, bold=True, size=9)
        _set_cell(table.rows[i].cells[1], hint, size=9)

    # Recorder signatures
    section_num += 1
    doc.add_heading(f"{section_num}. 기록자 서명", level=1)
    sig_points = followup.get("visits", ["Baseline", "Procedure", "F/U"])
    table = doc.add_table(rows=len(sig_points) + 1, cols=3)
    _format_table(table)
    for j, h in enumerate(["시점", "기록자", "서명/날짜"]):
        _set_cell(table.rows[0].cells[j], h, bold=True, size=8,
                  color=COLORS["section"])
    for i, visit in enumerate(sig_points):
        _set_cell(table.rows[i + 1].cells[0], visit, size=8)
        _set_cell(table.rows[i + 1].cells[1], "", size=8)
        _set_cell(table.rows[i + 1].cells[2], "", size=8)

    doc.save(output_path)
    return output_path


# ============================================================
# Main entry point
# ============================================================

def generate_research_docs(config, output_dir="."):
    """
    Generate all 3 research documents from a single config dict.

    Args:
        config: dict with study configuration (see EXAMPLE_CONFIG)
        output_dir: directory to save output files

    Returns:
        dict with paths: {"preexam": ..., "quicksheet": ..., "crf": ...}
    """
    os.makedirs(output_dir, exist_ok=True)
    study = config["study"]
    prefix = study.get("file_prefix", study["short_name"])

    paths = {}
    paths["preexam"] = _build_preexam(
        config,
        os.path.join(output_dir, f"{prefix}_예진지.docx"),
    )
    paths["quicksheet"] = _build_quicksheet(
        config,
        os.path.join(output_dir, f"{prefix}_Quick_Sheet.docx"),
    )
    paths["crf"] = _build_crf(
        config,
        os.path.join(output_dir, f"{prefix}_CRF.docx"),
    )

    print(f"Generated 3 documents in {output_dir}/:")
    for key, path in paths.items():
        print(f"  {key:12s}: {os.path.basename(path)}")

    return paths


# ============================================================
# Example config (FSHD study)
# ============================================================

EXAMPLE_CONFIG = {
    "study": {
        "short_name": "FSHD",
        "file_prefix": "FSHD",
        "title_ko": "동결견 환자의 초음파 유도 관절낭 보존 수압팽창술",
        "title_en": "US-guided capsule-preserving hydrodilatation for adhesive capsulitis",
        "department": "부산대학교병원 재활의학과",
        "pi": "이재현",
        "preexam_title": "어깨 통증 예진 문진표",
        "quicksheet_title": "FSHD Ultrasound Quick Sheet",
    },

    "subjects": {
        "fields": ["이름", "등록번호", "날짜"],
        "crf_fields": ["Subject ID", "병록번호(MRN)", "이니셜", "성명", "동의일"],
    },

    "enrollment": [
        {"label": "성별", "hint": "M(0) / F(1)", "source": "emr"},
        {"label": "생년월일 / 나이", "hint": "____.__.__  /  만 ___세", "source": "emr"},
        {"label": "이환측", "hint": "Rt(0) / Lt(1)", "source": "patient"},
        {"label": "우세손", "hint": "Rt(0) / Lt(1)", "source": "patient"},
        {"label": "키 (cm)", "hint": "___ cm", "source": "patient"},
        {"label": "몸무게 (kg)", "hint": "___ kg", "source": "patient"},
        {"label": "BMI", "hint": "자동계산", "source": "calc"},
    ],

    "inclusion": [
        "40–75세",
        "Active + Passive ROM 제한 ≥ 30° (2방향 이상)",
        "Primary adhesive capsulitis (이차성 원인 배제)",
        "증상 기간 ≥ 4주",
    ],

    "exclusion": [
        "이차성 원인 (골절, 수술, 석회건염 등)",
        "신경학적 결손 (상완신경총 등)",
        "최근 3개월 이내 스테로이드 주사",
        "전신 스테로이드 사용 중",
        "조절되지 않는 당뇨 (HbA1c > 9%)",
    ],

    "baseline": {
        "sections": [
            {
                "name": "동결견 병력",
                "items": [
                    {"label": "증상 시작일", "hint": "____년 ____월", "source": "patient"},
                    {"label": "이환 기간", "hint": "자동계산 (____개월)", "source": "calc"},
                    {"label": "Phase", "hint": "1 Freezing / 2 Frozen / 3 Thawing", "source": "examiner"},
                    {"label": "이전 치료 (주사 외)", "hint": "PT / 약물 / 도수 / 기타", "source": "patient"},
                    {"label": "이전 주사 횟수", "hint": "__회, 최종: ____년 ____월", "source": "patient"},
                    {"label": "당뇨 여부", "hint": "없음(0) / 있음(1)", "source": "patient"},
                ],
            },
            {
                "name": "Pain NRS (/10)",
                "items": [
                    {"label": "Resting pain", "hint": "___ /10", "source": "patient"},
                    {"label": "Activity pain", "hint": "___ /10", "source": "patient"},
                    {"label": "Night pain", "hint": "___ /10", "source": "patient"},
                ],
            },
            {
                "name": "Passive ROM (Goniometer)",
                "items": [
                    {"label": "Forward Flexion", "hint": "Affected ___° / Unaffected ___°", "source": "examiner"},
                    {"label": "Abduction", "hint": "Affected ___° / Unaffected ___°", "source": "examiner"},
                    {"label": "External Rotation", "hint": "Affected ___° / Unaffected ___°", "source": "examiner"},
                    {"label": "Extension", "hint": "Affected ___° / Unaffected ___°", "source": "examiner"},
                    {"label": "Internal Rotation", "hint": "Affected ___ / Unaffected ___", "source": "examiner"},
                ],
            },
            {
                "name": "SPADI + PGA",
                "items": [
                    {"label": "SPADI Pain (P1–P5)", "hint": "___ /50", "source": "patient"},
                    {"label": "SPADI Disability (D1–D8)", "hint": "___ /80", "source": "patient"},
                    {"label": "SPADI Total (%)", "hint": "자동계산", "source": "calc"},
                    {"label": "Patient Global Assessment", "hint": "___ /10", "source": "patient"},
                ],
            },
        ],
    },

    "procedure": {
        "sections": [
            {
                "name": "Ultrasound Baseline",
                "color": "FFFF99",
                "items": [
                    {"label": "CHL thickness (mm)", "hint": "Affected ___mm / Unaffected ___mm", "source": "examiner"},
                    {"label": "Axillary recess (mm)", "hint": "Affected ___mm / Unaffected ___mm", "source": "examiner"},
                    {"label": "SST status", "hint": "Intact / Partial / Full tear", "source": "examiner"},
                    {"label": "RI thickness (mm)", "hint": "Affected ___mm / Unaffected ___mm", "source": "examiner"},
                    {"label": "Power Doppler", "hint": "+  /  −", "source": "examiner"},
                    {"label": "Image file no.", "hint": "________", "source": "examiner"},
                ],
            },
            {
                "name": "Intra-Procedure",
                "color": "FFCCCC",
                "items": [
                    {"label": "Approach", "hint": "Anterior(0) / Posterior(1)", "source": "procedure"},
                    {"label": "5 mL", "hint": "Time / Resistance(0-3) / Discomfort / US", "source": "procedure"},
                    {"label": "10 mL", "hint": "Time / Resistance(0-3) / Discomfort / US", "source": "procedure"},
                    {"label": "15 mL", "hint": "Time / Resistance(0-3) / Discomfort / US", "source": "procedure"},
                    {"label": "20 mL", "hint": "Time / Resistance(0-3) / Discomfort / US", "source": "procedure"},
                    {"label": "25 mL", "hint": "Time / Resistance(0-3) / Discomfort / US", "source": "procedure"},
                    {"label": "30 mL", "hint": "Time / Resistance(0-3) / Discomfort / US", "source": "procedure"},
                    {"label": "__ mL (custom)", "hint": "Time / Resistance(0-3) / Discomfort / US", "source": "procedure"},
                ],
            },
            {
                "name": "Procedure End",
                "color": "FFCCCC",
                "items": [
                    {"label": "Total injected volume", "hint": "___ mL", "source": "procedure"},
                    {"label": "Total injection time", "hint": "___ sec", "source": "procedure"},
                    {"label": "Stopping reason", "hint": "1 Rupture / 2 Patient / 3 Resistance / 4 Plateau / 5 Target", "source": "procedure"},
                    {"label": "Capsule status", "hint": "Preserved(0) / Ruptured(1)", "source": "procedure"},
                    {"label": "Reinjection", "hint": "No(0) / Yes(1) → 추가 ___mL", "source": "procedure"},
                    {"label": "Rupture location", "hint": "Ant / Post / Inf / Sup / N/A", "source": "procedure"},
                ],
            },
            {
                "name": "AE Observation (15–30 min)",
                "color": "FFCCCC",
                "items": [
                    {"label": "Vasovagal reaction", "hint": "No(0) / Yes(1)", "source": "procedure"},
                    {"label": "Injection site pain", "hint": "None(0) / Mild(1) / Moderate(2) / Severe(3)", "source": "procedure"},
                    {"label": "Swelling / Ecchymosis", "hint": "No(0) / Yes(1)", "source": "procedure"},
                    {"label": "Other", "hint": "________________________", "source": "procedure"},
                ],
            },
        ],
    },

    "followup": {
        "visits": ["V1 Baseline", "V2 30min", "V3 2wk", "V4 4wk", "V5 8wk", "V6 12wk"],
        "measures": ["NRS (rest/act/night)", "ROM", "SPADI", "PGA", "AE check"],
    },

    "emr_post_entry": ["Subject ID", "성별", "생년월일", "시술일", "시술자"],
    "auto_calc": ["BMI", "이환기간", "SPADI Total %", "ROM ratio", "ROM change"],
}


# ============================================================
# CLI
# ============================================================

def main():
    parser = argparse.ArgumentParser(
        description="임상연구 문서 3종 (예진지, Quick Sheet, CRF) 자동 생성",
    )
    parser.add_argument(
        "config",
        nargs="?",
        help="Config file path (JSON or YAML). Omit with --example to print example config.",
    )
    parser.add_argument(
        "--output-dir", "-o",
        default=".",
        help="Output directory (default: current directory)",
    )
    parser.add_argument(
        "--example",
        action="store_true",
        help="Print example config (JSON) to stdout",
    )
    args = parser.parse_args()

    if args.example:
        print(json.dumps(EXAMPLE_CONFIG, ensure_ascii=False, indent=2))
        return

    if not args.config:
        parser.print_help()
        return

    # Load config
    config_path = args.config
    if config_path.endswith((".yaml", ".yml")):
        try:
            import yaml
            with open(config_path, "r", encoding="utf-8") as f:
                config = yaml.safe_load(f)
        except ImportError:
            print("ERROR: pyyaml required for YAML config. Install: pip install pyyaml")
            sys.exit(1)
    else:
        with open(config_path, "r", encoding="utf-8") as f:
            config = json.load(f)

    generate_research_docs(config, output_dir=args.output_dir)


if __name__ == "__main__":
    main()
